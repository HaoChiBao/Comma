"""Builds fixtures/WorldState_Demo.xlsx from the same rows the API seeds.

    pip install openpyxl
    PYTHONPATH=services/api python fixtures/build_workbook.py
"""

import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Font, PatternFill

from app.seed_data import WORKBOOK_ID, WORKBOOK_PERIOD, WORKBOOK_ROWS, WORKBOOK_SHEET

MONEY = '"$"#,##0.00'


def build(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKBOOK_SHEET
    sheet.append(["Vendor", "PO line 1", "PO line 2", "Accrual", "Period", "Ref"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row, vendor, _entity, line_1, line_2, manual, reference in WORKBOOK_ROWS:
        sheet.cell(row=row, column=1, value=vendor)
        if line_1 is not None:
            sheet.cell(row=row, column=2, value=float(Decimal(line_1)))
            sheet.cell(row=row, column=3, value=float(Decimal(line_2)))
            sheet.cell(row=row, column=4, value=f"=B{row}+C{row}")
        else:
            sheet.cell(row=row, column=4, value=float(Decimal(manual)))
        sheet.cell(row=row, column=5, value=WORKBOOK_PERIOD)
        sheet.cell(row=row, column=6, value=reference or None)
        for column in (2, 3, 4):
            sheet.cell(row=row, column=column).number_format = MONEY
    for column, width in zip("ABCDEF", (24, 12, 12, 12, 10, 10)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    # The add-in reads this to know which registered workbook it is looking at.
    workbook.custom_doc_props.append(StringProperty(name="WorldStateWorkbookId", value=WORKBOOK_ID))
    workbook.save(path)
    return path


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("WorldState_Demo.xlsx")
    print(build(target))
